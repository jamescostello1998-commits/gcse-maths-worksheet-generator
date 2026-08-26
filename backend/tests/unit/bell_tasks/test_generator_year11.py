import io

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.bell_tasks.generator import NUM_SLIDES, QUESTIONS_PER_TOPIC, YEAR_11, generate_bell_tasks_pptx
from app.bell_tasks.layout import GEOMETRY_YEAR_11, box_bounds
from app.core.errors import TopicNotFoundError

# A deliberate mix: some diagram-bearing topics, some not. Deliberately
# avoids any topic using the \colvec{}{} marker (vectors_arithmetic_H etc.) -
# a real, pre-existing gap (found via visual QA while building this test
# module, confirmed to affect the original Year 10 template too, not
# something this session introduced) where Bell Tasks' math_tokenizer/omml
# pipeline doesn't yet render a native column vector, unlike the PDF/DOCX
# renderers - left as-is, out of scope for the Year 11 restructuring.
EIGHT_TOPIC_IDS = [
    "angles_triangle_F",  # diagram
    "area_rectangle_F",  # diagram
    "fractions_add_subtract_F",  # no diagram, always produces a fraction
    "linear_two_step_F",  # no diagram
    "probability_single_event_F",  # diagram
    "bar_chart_construct_F",  # diagram, sometimes long prompt
    "algebraic_indices_H",  # exponent equation object
    "circle_parts_F",  # diagram, no verify()
]

GRID_TABLE_SHAPE_NAME = "Google Shape;146;p3"

_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _grid_table(slide):
    for shape in slide.shapes:
        if shape.has_table and shape.name == GRID_TABLE_SHAPE_NAME:
            return shape.table
    raise AssertionError("grid table not found on slide")


def _cell_full_text(cell) -> str:
    """See test_generator.py's identical helper - flattens a cell's
    paragraph back to text, including native-equation content read from each
    embedded equation's own mc:Fallback plain text."""
    parts = []
    for child in cell.text_frame.paragraphs[0]._p:
        tag = etree.QName(child).localname
        if tag == "r":
            texts = child.findall(f".//{{{_A_NS}}}t")
            parts.append("".join(t.text or "" for t in texts))
        elif tag == "AlternateContent":
            fallback_texts = child.findall(f".//{{{_MC_NS}}}Fallback//{{{_A_NS}}}t")
            parts.append("".join(t.text or "" for t in fallback_texts))
    return "".join(parts)


def test_generates_five_slides_with_eight_distinct_boxes_each():
    data = generate_bell_tasks_pptx(EIGHT_TOPIC_IDS, year=YEAR_11)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == NUM_SLIDES

    for slide in prs.slides:
        table = _grid_table(slide)
        for box in range(1, 9):
            row, col = GEOMETRY_YEAR_11.box_to_row_col[box]
            cell = table.cell(row, col)
            text = _cell_full_text(cell)
            assert text.startswith(f"{box}. ")
            assert len(text) > len(f"{box}. ")


def test_box_numbering_is_row_major_across_the_4_column_grid():
    # 1-4 on the top row (left to right), 5-8 on the bottom row - per direct
    # user confirmation, since this template's own cells ship blank (no
    # pre-printed numbers to match, unlike Year 10's column-major grid).
    assert GEOMETRY_YEAR_11.box_to_row_col == {
        1: (0, 0), 2: (0, 1), 3: (0, 2), 4: (0, 3),
        5: (1, 0), 6: (1, 1), 7: (1, 2), 8: (1, 3),
    }


def _box_picture_blobs(slide, box) -> tuple:
    bx, by, bw, bh = box_bounds(box, GEOMETRY_YEAR_11)
    blobs = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        if bx <= shape.left < bx + bw and by <= shape.top < by + bh:
            blobs.append(shape.image.blob)
    return tuple(blobs)


def test_each_box_shows_five_distinct_questions_across_the_week():
    data = generate_bell_tasks_pptx(EIGHT_TOPIC_IDS, year=YEAR_11)
    prs = Presentation(io.BytesIO(data))

    for box in range(1, 9):
        row, col = GEOMETRY_YEAR_11.box_to_row_col[box]
        seen = set()
        for slide in prs.slides:
            table = _grid_table(slide)
            cell = table.cell(row, col)
            text = _cell_full_text(cell)
            key = (text, _box_picture_blobs(slide, box))
            seen.add(key)
        assert len(seen) == QUESTIONS_PER_TOPIC, f"box {box} did not get {QUESTIONS_PER_TOPIC} distinct questions"


def test_runs_use_calibri_for_words_and_cambria_math_for_numbers():
    data = generate_bell_tasks_pptx(EIGHT_TOPIC_IDS, year=YEAR_11)
    prs = Presentation(io.BytesIO(data))

    seen_fonts = set()
    for slide in prs.slides:
        table = _grid_table(slide)
        for box in range(1, 9):
            row, col = GEOMETRY_YEAR_11.box_to_row_col[box]
            cell = table.cell(row, col)
            for run in cell.text_frame.paragraphs[0].runs:
                assert run.font.size is not None
                assert run.font.size.pt == 18
                seen_fonts.add(run.font.name)

    assert "Calibri" in seen_fonts
    assert "Cambria Math" in seen_fonts


def test_pictures_stay_within_their_own_cell_bounds():
    data = generate_bell_tasks_pptx(EIGHT_TOPIC_IDS, year=YEAR_11)
    prs = Presentation(io.BytesIO(data))

    for slide in prs.slides:
        for picture in (s for s in slide.shapes if s.shape_type == 13):
            matched = False
            for box in range(1, 9):
                cell_left, cell_top, cell_width, cell_height = box_bounds(box, GEOMETRY_YEAR_11)
                if (
                    picture.left >= cell_left
                    and picture.top >= cell_top
                    and picture.left + picture.width <= cell_left + cell_width
                    and picture.top + picture.height <= cell_top + cell_height
                ):
                    matched = True
                    break
            assert matched, f"picture at ({picture.left},{picture.top}) doesn't fit inside any cell"


def test_rejects_wrong_number_of_topics():
    with pytest.raises(ValueError):
        generate_bell_tasks_pptx(EIGHT_TOPIC_IDS[:6], year=YEAR_11)


def test_rejects_duplicate_topics():
    duplicated = EIGHT_TOPIC_IDS[:7] + [EIGHT_TOPIC_IDS[0]]
    with pytest.raises(ValueError):
        generate_bell_tasks_pptx(duplicated, year=YEAR_11)


def test_rejects_unknown_topic_id():
    bad = EIGHT_TOPIC_IDS[:7] + ["not_a_real_topic"]
    with pytest.raises(TopicNotFoundError):
        generate_bell_tasks_pptx(bad, year=YEAR_11)


def test_rejects_unknown_year():
    with pytest.raises(ValueError):
        generate_bell_tasks_pptx(EIGHT_TOPIC_IDS, year="year_12")


def test_year_10_six_topics_are_rejected_under_year_11():
    from tests.unit.bell_tasks.test_generator import SIX_TOPIC_IDS

    with pytest.raises(ValueError):
        generate_bell_tasks_pptx(SIX_TOPIC_IDS, year=YEAR_11)
